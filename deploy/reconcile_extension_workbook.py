"""Compare the supplied extension workbook with a read-only CSV snapshot of extensions."""

from collections import Counter
import csv
import json
from pathlib import Path
import sys

import openpyxl


COLUMNS = ("building", "extension_number", "floor", "department", "remarks", "did_dod")
LIMITS = (20, 50, 50, 100, 255, 100)
HEADERS = ("층수", "부서", "비고", "DID, DOD")


def cell_text(value):
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("Boolean values are not valid extension data")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def workbook_rows(path):
    # The source has an out-of-range style reference; streaming mode reads its values safely.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(workbook.worksheets) != 1:
        raise ValueError("Expected exactly one worksheet")
    sheet = workbook.active
    result = []
    sections = []
    building = None
    errors = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = row[:5]
        if not any(cell_text(value) is not None for value in values):
            if any(cell_text(value) is not None for value in row[5:]):
                errors.append(f"Data outside columns A:E on row {row_number}")
            continue
        if tuple(cell_text(value) for value in values[1:]) == HEADERS:
            label = cell_text(values[0])
            building = "0" if label == "원외" else label
            sections.append((row_number, building))
            continue
        if building is None:
            errors.append(f"Data row before a building header: {row_number}")
            continue
        mapped = (building,) + tuple(cell_text(value) for value in values)
        # The external pediatric contact is a named contact, not an extension number.
        if (building == "0" and mapped[1] and not mapped[1].isdigit()
                and mapped[2:5] == (None, None, None) and mapped[5]):
            mapped = (building, None, None, None, mapped[1], mapped[5])
        for field, value, limit in zip(COLUMNS, mapped, LIMITS):
            if value is not None and len(value) > limit:
                errors.append(f"Row {row_number}: {field} exceeds {limit} characters")
        if mapped[1] is None and mapped[4] is None:
            errors.append(f"Row {row_number}: extension number is blank")
        if mapped[1] is not None and not mapped[1].isdigit():
            errors.append(f"Row {row_number}: extension number is not numeric")
        result.append((row_number, mapped))
    formula_sheet = openpyxl.load_workbook(path, read_only=True, data_only=False).active
    formulas = [(cell.coordinate, cell.value) for row in formula_sheet.iter_rows(min_col=1, max_col=5)
                for cell in row if cell.data_type == "f"]
    if formulas:
        errors.append(f"Formula cells in A:E: {len(formulas)}")
    return result, sections, errors


def database_rows(path):
    with open(path, newline="", encoding="utf-8-sig") as source:
        reader = csv.DictReader(source)
        if reader.fieldnames != ["id", *COLUMNS]:
            raise ValueError(f"Unexpected database columns: {reader.fieldnames}")
        return [(int(row["id"]), tuple(cell_text(row[column]) for column in COLUMNS)) for row in reader]


def compare(workbook_path, db_path):
    excel, sections, errors = workbook_rows(workbook_path)
    database = database_rows(db_path)
    excel_values = [row for _, row in excel]
    db_values = [row for _, row in database]
    by_id_equal = sum(left == right for left, right in zip(excel_values, db_values))
    changes = Counter()
    for left, right in zip(excel_values, db_values):
        for column, new_value, old_value in zip(COLUMNS, left, right):
            if new_value != old_value:
                changes[column] += 1
    excel_numbers = Counter(row[1] for row in excel_values)
    db_numbers = Counter(row[1] for row in db_values)
    summary = {
        "excel_records": len(excel),
        "database_records": len(database),
        "database_ids_sequential": [key for key, _ in database] == list(range(1, len(database) + 1)),
        "sections": sections,
        "excel_buildings": dict(Counter(row[0] for row in excel_values)),
        "database_buildings": dict(Counter(row[0] for row in db_values)),
        "equal_by_position": by_id_equal,
        "equal_full_row_multiset": sum((Counter(excel_values) & Counter(db_values)).values()),
        "same_extension_multiset": excel_numbers == db_numbers,
        "new_extension_count": sum((excel_numbers - db_numbers).values()),
        "removed_extension_count": sum((db_numbers - excel_numbers).values()),
        "changed_fields_by_position": dict(changes),
        "errors": errors,
    }
    return summary, excel


if __name__ == "__main__":
    summary, excel_rows = compare(Path(sys.argv[1]), Path(sys.argv[2]))
    print(json.dumps(summary, ensure_ascii=True))
    if len(sys.argv) > 3:
        if summary["errors"] or summary["excel_records"] != summary["database_records"]:
            raise SystemExit("Validation failed; import CSV was not written")
        output_path = Path(sys.argv[3])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", newline="", encoding="utf-8") as destination:
            writer = csv.writer(destination)
            writer.writerow(("id", *COLUMNS))
            for record_id, (_, fields) in enumerate(excel_rows, 1):
                writer.writerow((record_id, *(value if value is not None else "" for value in fields)))
